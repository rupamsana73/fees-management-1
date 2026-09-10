import csv
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import DecimalField, Q, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render, reverse
from django.utils.html import format_html
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from students.decorators import admin_required
from students.models import Student
from users.models import User

from .forms import FeePaymentForm
from .models import FeePayment


@login_required
@admin_required
def payment_add(request):
    form = FeePaymentForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        payment = form.save()
        messages.success(
            request,
            format_html(
                'Payment added successfully. <a href="{}">View Receipt</a> | <a href="{}">Download PDF</a>',
                reverse("payment-receipt", args=[payment.pk]),
                reverse("payment-receipt-pdf", args=[payment.pk]),
            ),
        )
        return redirect("student-list")

    students = Student.objects.annotate(
        paid_total=Coalesce(
            Sum("payments__amount_paid"),
            Value(Decimal("0.00")),
            output_field=DecimalField(max_digits=12, decimal_places=2),
        )
    )
    for student in students:
        student.pending_total = student.total_fee - student.paid_total
    return render(request, "payments/payment_form.html", {"form": form, "students": students})


@login_required
def payment_receipt(request, payment_id):
    payment = get_object_or_404(FeePayment, pk=payment_id)
    if request.user.role == User.Role.ADMIN:
        allowed = True
    elif request.user.role == User.Role.STUDENT:
        allowed = payment.student.user_id == request.user.id
    else:
        allowed = False

    if not allowed:
        return HttpResponse("Forbidden", status=403)

    context = build_receipt_context(payment)
    context["student"] = payment.student
    context["base_template"] = "student_base.html" if request.user.role == User.Role.STUDENT else "admin_base.html"
    context["back_url"] = (
        reverse("payment-history") if request.user.role == User.Role.STUDENT else reverse("student-list")
    )
    return render(request, "payments/receipt.html", context)


@login_required
def payment_receipt_pdf(request, payment_id):
    payment = get_object_or_404(FeePayment, pk=payment_id)
    if request.user.role == User.Role.ADMIN:
        allowed = True
    elif request.user.role == User.Role.STUDENT:
        allowed = payment.student.user_id == request.user.id
    else:
        allowed = False

    if not allowed:
        return HttpResponse("Forbidden", status=403)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("FeeFlow", styles["Title"]))
    story.append(Paragraph("STUDENT FEE MANAGEMENT SYSTEM", styles["Heading2"]))
    story.append(Paragraph("Debasis Kamila", styles["Heading3"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph("PAYMENT RECEIPT", styles["Heading1"]))
    story.append(Spacer(1, 8))
    data = [
        ("Receipt No:", payment.receipt_number),
        ("Payment Date:", payment.payment_date.strftime("%d %B %Y")),
        ("Student Name:", payment.student.name),
        ("Student ID:", payment.student.student_id),
        ("Course:", payment.student.course),
        ("Email:", payment.student.email),
        ("Phone:", payment.student.phone),
        ("Amount Paid:", f"₹{payment.amount_paid:,.2f}"),
        ("Payment Method:", payment.payment_method or "Cash"),
        ("Transaction ID:", payment.transaction_id or "—"),
        ("Remarks:", payment.remarks or "—"),
        ("Total Fee:", f"₹{payment.student.total_fee:,.2f}"),
        ("Previously Paid:", f"₹{payment.previous_paid:,.2f}"),
        ("This Payment:", f"₹{payment.amount_paid:,.2f}"),
        ("Total Paid:", f"₹{payment.total_paid_after_payment:,.2f}"),
        ("Remaining Balance:", f"₹{payment.remaining_balance:,.2f}"),
        ("Payment Status:", payment.payment_status),
    ]
    for label, value in data:
        story.append(Paragraph(f"<b>{label}</b> {value}", styles["BodyText"]))
    story.append(Spacer(1, 18))
    story.append(Paragraph("Thank you.", styles["BodyText"]))
    story.append(Paragraph("Debasis Kamila", styles["BodyText"]))
    story.append(Paragraph("FeeFlow Student Fee Management System", styles["BodyText"]))
    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="FeeFlow_Receipt_{payment.receipt_number}.pdf"'
    )
    return response


def build_receipt_context(payment):
    previous_paid = payment.previous_paid
    total_paid_after = payment.total_paid_after_payment
    pending = payment.remaining_balance
    return {
        "payment": payment,
        "receipt_number": payment.receipt_number,
        "previous_paid": previous_paid,
        "current_payment": payment.amount_paid,
        "total_paid_after_payment": total_paid_after,
        "pending": pending,
        "payment_status": payment.payment_status,
    }


@login_required
@admin_required
def reports(request):
    context = build_report_context(request)
    return render(
        request,
        "payments/reports.html",
        context,
    )


def build_report_context(request):
    selected_student_id = request.GET.get("student", "")
    selected_course = request.GET.get("course", "")
    selected_payment_status = request.GET.get("payment_status", "")
    selected_payment_method = request.GET.get("payment_method", "")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")

    filters = Q()
    if selected_student_id:
        filters &= Q(student_id=selected_student_id)
    if selected_course:
        filters &= Q(student__course=selected_course)
    if selected_payment_method:
        filters &= Q(payment_method=selected_payment_method)
    if date_from:
        filters &= Q(payment_date__gte=date_from)
    if date_to:
        filters &= Q(payment_date__lte=date_to)

    payment_list = list(FeePayment.objects.select_related("student").filter(filters).order_by("-payment_date", "-pk"))
    if selected_payment_status:
        payment_list = [
            payment for payment in payment_list
            if payment.payment_status == selected_payment_status.upper()
        ]

    students = Student.objects.filter(
        **({"pk": selected_student_id} if selected_student_id else {}),
        **({"course": selected_course} if selected_course else {}),
    ).order_by("name")
    payments_by_student = {}
    for payment in payment_list:
        payments_by_student.setdefault(payment.student_id, []).append(payment)

    student_rows = []
    for student in students:
        student_payments = payments_by_student.get(student.pk, [])
        paid_total = sum((payment.amount_paid for payment in student_payments), Decimal("0.00"))
        student_rows.append({
            "id": student.pk,
            "name": student.name,
            "student_id": student.student_id,
            "course": student.course,
            "total_fee": student.total_fee,
            "paid_total": paid_total,
            "pending_total": max(student.total_fee - paid_total, Decimal("0.00")),
            "last_payment_date": student_payments[0].payment_date if student_payments else None,
        })

    total_fees = sum((row["total_fee"] for row in student_rows), Decimal("0.00"))
    total_collected = sum((payment.amount_paid for payment in payment_list), Decimal("0.00"))
    status_rows = [
        {"label": "Paid", "total": sum(payment.payment_status == "PAID" for payment in payment_list)},
        {"label": "Partial", "total": sum(payment.payment_status == "PARTIAL" for payment in payment_list)},
        {"label": "Unpaid", "total": sum(payment.payment_status == "PENDING" for payment in payment_list)},
    ]

    monthly_totals = {}
    course_totals = {}
    method_totals = {}
    for payment in payment_list:
        monthly_totals[payment.payment_date.strftime("%B %Y")] = monthly_totals.get(
            payment.payment_date.strftime("%B %Y"), Decimal("0.00")
        ) + payment.amount_paid
        course = payment.student.course or "Unassigned"
        method = payment.payment_method or "Cash"
        course_totals[course] = course_totals.get(course, Decimal("0.00")) + payment.amount_paid
        method_totals[method] = method_totals.get(method, Decimal("0.00")) + payment.amount_paid

    payment_rows = [{
        "id": payment.pk,
        "receipt_number": payment.receipt_number,
        "payment_date": payment.payment_date,
        "student_name": payment.student.name,
        "student_id": payment.student.student_id,
        "course": payment.student.course,
        "amount_paid": payment.amount_paid,
        "payment_method": payment.payment_method or "Cash",
        "transaction_id": payment.transaction_id or "",
        "status": payment.payment_status,
        "student_url": reverse("payment-receipt", args=[payment.pk]),
    } for payment in payment_list]

    return {
        "student_rows": student_rows,
        "payment_rows": payment_rows,
        "report_total_students": len(student_rows),
        "report_total_fees": total_fees,
        "report_collected": total_collected,
        "report_pending": max(total_fees - total_collected, Decimal("0.00")),
        "report_fully_paid": sum(row["paid_total"] >= row["total_fee"] for row in student_rows),
        "report_partially_paid": sum(Decimal("0.00") < row["paid_total"] < row["total_fee"] for row in student_rows),
        "report_unpaid": sum(row["paid_total"] == Decimal("0.00") for row in student_rows),
        "monthly_collections": [{"label": label, "total": total} for label, total in monthly_totals.items()],
        "course_collections": [{"label": label, "total": total} for label, total in sorted(course_totals.items(), key=lambda item: item[1], reverse=True)],
        "method_collections": [{"label": label, "total": total} for label, total in sorted(method_totals.items(), key=lambda item: item[1], reverse=True)],
        "status_rows": status_rows,
        "students_for_dropdown": Student.objects.order_by("name"),
        "course_options": Student.objects.exclude(course="").order_by("course").values_list("course", flat=True).distinct(),
        "payment_method_options": FeePayment.objects.exclude(payment_method="").order_by("payment_method").values_list("payment_method", flat=True).distinct(),
        "selected_student": selected_student_id,
        "selected_course": selected_course,
        "selected_payment_status": selected_payment_status,
        "selected_payment_method": selected_payment_method,
        "date_from": date_from,
        "date_to": date_to,
        "applied_filters": [
            ("Date From", date_from or "All"),
            ("Date To", date_to or "All"),
            ("Student", next((row["name"] for row in student_rows if str(row["id"]) == selected_student_id), "All")),
            ("Course", selected_course or "All"),
            ("Payment Status", selected_payment_status or "All"),
            ("Payment Method", selected_payment_method or "All"),
        ],
    }


def _export_response(content, content_type, filename):
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
@admin_required
def reports_csv(request):
    context = build_report_context(request)
    rows = context["payment_rows"]
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="FeeFlow_Report.csv"'
    writer = csv.writer(response)
    writer.writerow(["Receipt Number", "Payment Date", "Student Name", "Student ID", "Course", "Amount Paid", "Payment Method", "Transaction Reference", "Status"])
    for row in rows:
        writer.writerow([row["receipt_number"], row["payment_date"].isoformat(), row["student_name"], row["student_id"], row["course"], f'{row["amount_paid"]:.2f}', row["payment_method"], row["transaction_id"], row["status"]])
    return response


@login_required
@admin_required
def reports_excel(request):
    context = build_report_context(request)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["FeeFlow Financial Report"])
    summary.append(["Generated", date.today()])
    summary.append([])
    summary.append(["Metric", "Value"])
    for label, value in (("Total Students", context["report_total_students"]), ("Total Fees", context["report_total_fees"]), ("Total Collected", context["report_collected"]), ("Total Pending", context["report_pending"]), ("Fully Paid", context["report_fully_paid"]), ("Partially Paid", context["report_partially_paid"]), ("Unpaid", context["report_unpaid"])):
        summary.append([label, value])
    summary.append([])
    summary.append(["Applied Filter", "Value"])
    for label, value in context["applied_filters"]:
        summary.append([label, value])
    _style_sheet(summary, 2, header_rows=(4, 13))

    student_sheet = workbook.create_sheet("Student Fee Report")
    student_sheet.append(["Student", "Student ID", "Course", "Total Fee", "Collected", "Pending", "Last Payment"])
    for row in context["student_rows"]:
        student_sheet.append([row["name"], row["student_id"], row["course"], row["total_fee"], row["paid_total"], row["pending_total"], row["last_payment_date"]])
    _style_sheet(student_sheet, 7, money_columns=(4, 5, 6), date_columns=(7,))

    payment_sheet = workbook.create_sheet("Payment Report")
    payment_sheet.append(["Receipt Number", "Date", "Student", "Student ID", "Course", "Amount", "Method", "Transaction Reference", "Status"])
    for row in context["payment_rows"]:
        payment_sheet.append([row["receipt_number"], row["payment_date"], row["student_name"], row["student_id"], row["course"], row["amount_paid"], row["payment_method"], row["transaction_id"], row["status"]])
    _style_sheet(payment_sheet, 9, money_columns=(6,), date_columns=(2,))

    for title, rows in (("Course Analytics", context["course_collections"]), ("Payment Method Analytics", context["method_collections"])):
        sheet = workbook.create_sheet(title)
        sheet.append(["Category", "Collected"])
        for row in rows:
            sheet.append([row["label"], row["total"]])
        _style_sheet(sheet, 2, money_columns=(2,))

    output = BytesIO()
    workbook.save(output)
    return _export_response(output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "FeeFlow_Report.xlsx")


def _style_sheet(sheet, column_count, money_columns=(), date_columns=(), header_rows=(1,)):
    header_fill = PatternFill("solid", fgColor="17324D")
    for row_number in header_rows:
        for cell in sheet[row_number]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2" if header_rows == (1,) else "A5"
    for column in money_columns:
        for cell in list(sheet.columns)[column - 1][1:]:
            cell.number_format = '#,##0.00'
    for column in date_columns:
        for cell in list(sheet.columns)[column - 1][1:]:
            cell.number_format = "dd mmm yyyy"
    for index in range(1, column_count + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 20


@login_required
@admin_required
def reports_pdf(request):
    context = build_report_context(request)
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=12 * mm, leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph("FeeFlow", styles["Title"]), Paragraph("Student Fee Management System", styles["Heading2"]), Paragraph("Financial Report", styles["Heading1"]), Paragraph(f"Generated: {date.today().strftime('%d %B %Y')}", styles["BodyText"]), Spacer(1, 8)]
    story.append(Paragraph("Applied Filters", styles["Heading3"]))
    story.append(Paragraph(" | ".join(f"{label}: {value}" for label, value in context["applied_filters"]), styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Summary", styles["Heading3"]))
    story.append(Table([["Total Students", "Total Fees", "Collected", "Pending"], [context["report_total_students"], f'{context["report_total_fees"]:.2f}', f'{context["report_collected"]:.2f}', f'{context["report_pending"]:.2f}']], repeatRows=1, style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), "#17324D"), ("TEXTCOLOR", (0, 0), (-1, 0), "white"), ("GRID", (0, 0), (-1, -1), 0.25, "#CCCCCC"), ("ALIGN", (0, 0), (-1, -1), "CENTER")])) )
    story.append(Spacer(1, 10))
    story.append(Paragraph("Student Fee Report", styles["Heading3"]))
    student_data = [["Student", "ID", "Course", "Total", "Collected", "Pending"]]
    student_data.extend([[row["name"], row["student_id"], row["course"], f'{row["total_fee"]:.2f}', f'{row["paid_total"]:.2f}', f'{row["pending_total"]:.2f}'] for row in context["student_rows"]])
    story.append(Table(student_data, repeatRows=1, style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), "#17324D"), ("TEXTCOLOR", (0, 0), (-1, 0), "white"), ("GRID", (0, 0), (-1, -1), 0.25, "#CCCCCC"), ("FONTSIZE", (0, 0), (-1, -1), 7)])))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Payment Report", styles["Heading3"]))
    payment_data = [["Receipt", "Date", "Student", "Course", "Amount", "Method", "Reference"]]
    payment_data.extend([[row["receipt_number"], row["payment_date"].strftime("%d %b %Y"), row["student_name"], row["course"], f'{row["amount_paid"]:.2f}', row["payment_method"], row["transaction_id"]] for row in context["payment_rows"]])
    story.append(Table(payment_data, repeatRows=1, style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), "#17324D"), ("TEXTCOLOR", (0, 0), (-1, 0), "white"), ("GRID", (0, 0), (-1, -1), 0.25, "#CCCCCC"), ("FONTSIZE", (0, 0), (-1, -1), 7)])))
    document.build(story)
    return _export_response(buffer.getvalue(), "application/pdf", "FeeFlow_Report.pdf")


@login_required
@admin_required
def reports_print(request):
    context = build_report_context(request)
    context["print_view"] = True
    context["today"] = date.today()
    return render(request, "payments/reports_print.html", context)
