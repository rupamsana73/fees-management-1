from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse

from payments.models import FeePayment
from students.models import Student
from students.utils import create_student_with_account
from users.models import User


class PaymentReceiptTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="teacher",
            password="TeacherPass123!",
            role=User.Role.ADMIN,
        )
        self.student, _, self.password = create_student_with_account(
            {
                "name": "Alice Student",
                "email": "alice@example.com",
                "course": "Computer Science",
                "total_fee": "1200.00",
                "phone": "9876543210",
                "admission_date": "2026-09-10",
                "status": "active",
            }
        )
        self.payment = FeePayment.objects.create(
            student=self.student,
            amount_paid=Decimal("300.00"),
            month="September 2026",
            remarks="Initial payment",
        )

    def test_admin_can_view_receipt(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("payment-receipt", args=[self.payment.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "PAYMENT RECEIPT")
        self.assertContains(response, self.student.name)

    def test_student_can_view_own_receipt(self):
        self.client.force_login(self.student.user)
        response = self.client.get(reverse("payment-receipt", args=[self.payment.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.student.student_id)

    def test_student_cannot_view_another_student_receipt(self):
        other_student, _, _ = create_student_with_account(
            {
                "name": "Bob Learner",
                "email": "bob@example.com",
                "course": "Physics",
                "total_fee": "800.00",
                "phone": "1234567890",
                "admission_date": "2026-09-11",
                "status": "active",
            }
        )
        other_payment = FeePayment.objects.create(
            student=other_student,
            amount_paid=Decimal("200.00"),
            month="September 2026",
            remarks="Other payment",
        )

        self.client.force_login(self.student.user)
        response = self.client.get(reverse("payment-receipt", args=[other_payment.pk]))

        self.assertIn(response.status_code, (302, 403, 404))

    def test_unauthenticated_user_cannot_access_receipt(self):
        response = self.client.get(reverse("payment-receipt", args=[self.payment.pk]))

        self.assertEqual(response.status_code, 302)
        self.assertIn("/login/", response.url)

    def test_receipt_contains_expected_payment_information(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("payment-receipt", args=[self.payment.pk]))

        self.assertContains(response, "FeeFlow")
        self.assertContains(response, self.student.name)
        self.assertContains(response, self.student.student_id)
        self.assertContains(response, self.student.course)
        self.assertContains(response, self.student.email)
        self.assertContains(response, str(self.payment.amount_paid))
        self.assertContains(response, self.payment.month)
        self.assertContains(response, "Initial payment")

    def test_receipt_number_is_generated_and_stable_after_edit(self):
        self.client.force_login(self.admin)
        first_response = self.client.get(reverse("payment-receipt", args=[self.payment.pk]))
        first_number = self.payment.receipt_number

        self.assertTrue(first_number)
        self.assertIn(first_number, first_response.content.decode())

        self.payment.amount_paid = Decimal("450.00")
        self.payment.save(update_fields=["amount_paid"])

        second_response = self.client.get(reverse("payment-receipt", args=[self.payment.pk]))
        self.assertIn(first_number, second_response.content.decode())
        self.assertEqual(self.payment.receipt_number, first_number)

    def test_receipt_pdf_response_has_correct_content_type(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("payment-receipt-pdf", args=[self.payment.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content)
        self.assertIn("FeeFlow_Receipt_", response["Content-Disposition"])
        self.assertIn(self.payment.receipt_number, response["Content-Disposition"])

    def test_admin_receipt_link_is_available_on_student_history_page(self):
        self.student.user.must_change_password = False
        self.student.user.save(update_fields=["must_change_password"])
        self.client.force_login(self.student.user)
        response = self.client.get(reverse("payment-history"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("payment-receipt", args=[self.payment.pk]))


class AdvancedReportsTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="teacher2",
            password="TeacherPass123!",
            role=User.Role.ADMIN,
        )
        self.student_user = User.objects.create_user(
            username="studentreport",
            password="StudentPass123!",
            role=User.Role.STUDENT,
        )
        self.student = Student.objects.create(
            user=self.student_user,
            student_id="STU200",
            name="Charlie Report",
            email="charlie@example.com",
            phone="9998887777",
            course="Computer Science",
            total_fee=Decimal("1500.00"),
            admission_date="2026-09-01",
            status=Student.Status.ACTIVE,
        )
        self.payment = FeePayment.objects.create(
            student=self.student,
            amount_paid=Decimal("600.00"),
            payment_date="2026-09-05",
            month="September 2026",
            payment_method="UPI",
            transaction_id="TXN-1001",
            remarks="Initial payment",
        )

    def test_reports_page_requires_admin(self):
        self.client.force_login(self.student_user)
        response = self.client.get(reverse("reports"))
        self.assertEqual(response.status_code, 302)

        self.client.force_login(self.admin)
        response = self.client.get(reverse("reports"))
        self.assertEqual(response.status_code, 200)

    def test_reports_page_displays_summary_cards_and_filters(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("reports"))

        self.assertContains(response, "Total Students")
        self.assertContains(response, "Total Fees")
        self.assertContains(response, "Total Collected")
        self.assertContains(response, "Total Pending")
        self.assertContains(response, "Date From")
        self.assertContains(response, "Date To")
        self.assertContains(response, "Payment Status")
        self.assertContains(response, "Student")
        self.assertContains(response, "Course")
        self.assertContains(response, "Payment Method")
        self.assertContains(response, "Student-wise Fee Report")
        self.assertContains(response, "Payment-wise Report")

    def test_reports_filters_update_report_values(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("reports"),
            {
                "student": str(self.student.pk),
                "course": "Computer Science",
                "payment_method": "UPI",
            },
        )

        self.assertContains(response, "Charlie Report")
        self.assertContains(response, "TXN-1001")
        self.assertContains(response, "UPI")
        self.assertContains(response, "Computer Science")

    def test_csv_export_contains_filtered_payment_data(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("reports-csv"),
            {"student": self.student.pk, "course": "Computer Science", "payment_method": "UPI"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment; filename=\"FeeFlow_Report.csv\"", response["Content-Disposition"])
        self.assertIn("Receipt Number,Payment Date,Student Name", response.content.decode("utf-8"))
        self.assertIn("Charlie Report", response.content.decode("utf-8"))
        self.assertIn("TXN-1001", response.content.decode("utf-8"))

    def test_excel_export_contains_report_sheets(self):
        from openpyxl import load_workbook

        self.client.force_login(self.admin)
        response = self.client.get(reverse("reports-excel"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(filename=BytesIO(response.content), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Student Fee Report", "Payment Report", "Course Analytics", "Payment Method Analytics"],
        )
        self.assertEqual(workbook["Payment Report"]["A2"].value, self.payment.receipt_number)

    def test_pdf_export_returns_report_pdf(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("reports-pdf"), {"course": "Computer Science"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment; filename=\"FeeFlow_Report.pdf\"", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_print_view_contains_filters_and_filtered_tables(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("reports-print"),
            {"student": self.student.pk, "course": "Computer Science", "payment_method": "UPI"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "FeeFlow")
        self.assertContains(response, "Computer Science")
        self.assertContains(response, "Charlie Report")
        self.assertContains(response, "TXN-1001")

    def test_students_and_unauthenticated_users_cannot_export_reports(self):
        for route_name in ("reports-csv", "reports-excel", "reports-pdf", "reports-print"):
            response = self.client.get(reverse(route_name))
            self.assertEqual(response.status_code, 302)
            self.assertIn("/login/", response.url)

            self.client.force_login(self.student_user)
            response = self.client.get(reverse(route_name))
            self.assertEqual(response.status_code, 302)
            self.assertIn("student-dashboard", response.url)
            self.client.logout()

    def test_combined_filters_and_empty_filtered_results(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("reports-csv"),
            {
                "date_from": "2099-01-01",
                "date_to": "2099-12-31",
                "student": self.student.pk,
                "course": "Computer Science",
                "payment_status": "PAID",
                "payment_method": "UPI",
            },
        )

        content = response.content.decode("utf-8")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Receipt Number,Payment Date,Student Name", content)
        self.assertNotIn("Charlie Report", content)
        self.assertNotIn("TXN-1001", content)
